#!/usr/bin/env python3
"""
Scrape the IND "Public register Work" (recognised sponsors) table into Excel.

Usage:
    pip install requests beautifulsoup4 lxml openpyxl
    python scrape_ind_sponsors.py                  # downloads the live page
    python scrape_ind_sponsors.py saved_page.html  # or parse a page saved from your browser

Output: ind_recognised_sponsors_work.xlsx  (columns: Organisation, KVK number)
"""
import sys
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

URL = "https://ind.nl/en/public-register-recognised-sponsors/public-register-work"
OUT = "ind_recognised_sponsors_work.xlsx"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "en",
}


def load_html(source=None):
    if source:
        return Path(source).read_text(encoding="utf-8")
    resp = requests.get(URL, headers=HEADERS, timeout=60)
    print(f"[fetch] HTTP {resp.status_code}  {len(resp.content):,} bytes  {resp.headers.get('content-type','')}")
    resp.raise_for_status()
    resp.encoding = "utf-8"
    return resp.text


def diagnose(html):
    """
    Report what the page actually contains, on every run.

    The three ways this scrape can fail all look the same from the outside — a
    small row count — so the run says which one happened rather than leaving it
    to be guessed at. A blocked request, a table rendered client-side, and a
    paginated table are each visible here.
    """
    soup = BeautifulSoup(html, "lxml")
    tables = soup.find_all("table")
    print(f"[diag] page length {len(html):,} chars, {len(tables)} table(s)")
    for i, t in enumerate(tables):
        header = [th.get_text(strip=True) for th in t.find_all("th")][:4]
        body = [tr for tr in t.find_all("tr") if tr.find_all("td")]
        print(f"[diag]   table {i}: {len(body):,} body rows, header={header}")

    lowered = html.lower()
    for marker in ("pagination", "data-page", 'rel="next"', "load more", "showing 1", "paginator"):
        if marker in lowered:
            print(f"[diag] possible pagination marker present: {marker!r}")
    if len(tables) == 0:
        print("[diag] NO <table> in the HTML — the register is probably rendered "
              "client-side, so requests+BeautifulSoup cannot see it.")


def parse_rows(html):
    """
    Read (Organisation, KVK number) out of the register table.

    Cells are collected from th *and* td, which is the whole trick. IND marks
    the first column of every data row as <th scope="row"> — correct HTML for a
    table whose rows are keyed by name — so looking only at td finds a single
    cell per row, fails the two-column check, and silently returns nothing. The
    live page has 12,984 rows and that mistake yielded 0.

    The header row is then skipped by its content rather than its position,
    since it is no longer distinguishable by tag.
    """
    soup = BeautifulSoup(html, "lxml")
    target = None
    for table in soup.find_all("table"):
        header = [th.get_text(strip=True) for th in table.find_all("th")]
        if any("organisation" in h.lower() for h in header):
            target = table
            break
    if target is None:
        raise RuntimeError("Could not find the Organisation table on the page.")

    rows = []
    for tr in target.find_all("tr"):
        cells = [c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
        if len(cells) < 2 or not cells[0]:
            continue
        if cells[0].strip().lower() == "organisation":
            continue  # the header row
        rows.append((cells[0], cells[1]))  # (Organisation, KVK number)
    return rows


def write_excel(rows, path=OUT):
    wb = Workbook()
    ws = wb.active
    ws.title = "Recognised sponsors - Work"

    ws.append(["Organisation", "KVK number"])
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(vertical="center")

    for org, kvk in rows:
        ws.append([org, kvk])

    # Keep KVK numbers as text so leading zeros (e.g. 02076358) are preserved
    for row in ws.iter_rows(min_row=2):
        row[0].font = Font(name="Arial")
        row[1].font = Font(name="Arial")
        row[1].number_format = "@"

    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 14
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else None
    html = load_html(src)
    diagnose(html)

    rows = parse_rows(html)
    write_excel(rows)

    # Duplicate organisation names are expected — the same name can appear with
    # different KVK numbers — so report both counts rather than one.
    unique_names = len({org for org, _ in rows})
    unique_kvk = len({kvk for _, kvk in rows})
    print(f"\nSaved {len(rows):,} rows to {OUT}")
    print(f"  distinct organisation names: {unique_names:,}")
    print(f"  distinct KVK numbers       : {unique_kvk:,}")
    print(f"  rows with an empty KVK     : {sum(1 for _, k in rows if not k):,}")
    print(f"  KVK values keeping a leading zero: {sum(1 for _, k in rows if k.startswith('0')):,}")
    print(f"First: {rows[0]}")
    print(f"Last:  {rows[-1]}")
